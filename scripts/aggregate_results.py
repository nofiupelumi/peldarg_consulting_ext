#!/usr/bin/env python3
import os
import json
import glob
import pandas as pd
from typing import Dict, List
import hmac
import hashlib
import requests


def norm_space(s: str) -> str:
    return " ".join((s or "").split()).strip()


def save_outputs(df: pd.DataFrame, base: str, out_dir: str, *, skip_docx: bool = False, max_docx_rows: int = 3000) -> Dict[str,str]:
    os.makedirs(out_dir, exist_ok=True)
    paths: Dict[str, str] = {}
    csv_path = os.path.join(out_dir, base + '.csv')
    xlsx_path = os.path.join(out_dir, base + '.xlsx')
    df.to_csv(csv_path, index=False)
    print(f"[agg] Wrote CSV: {csv_path}")
    try:
        df.to_excel(xlsx_path, index=False)
        paths['xlsx'] = xlsx_path
        print(f"[agg] Wrote XLSX: {xlsx_path}")
    except Exception as e:
        print(f"[agg] XLSX write failed: {e}")
    paths['csv'] = csv_path
    # Optionally write DOCX (can be very slow for large tables)
    if not skip_docx and len(df) <= max_docx_rows:
        try:
            from docx import Document  # lazy import
            docx_path = os.path.join(out_dir, base + '.docx')
            doc = Document()
            table = doc.add_table(rows=len(df)+1, cols=len(df.columns))
            for j, col in enumerate(df.columns):
                table.cell(0,j).text = col
            for i, row in enumerate(df.itertuples(index=False), start=1):
                for j, val in enumerate(row):
                    table.cell(i,j).text = '' if pd.isna(val) else str(val)
            doc.save(docx_path)
            paths['docx'] = docx_path
            print(f"[agg] Wrote DOCX: {docx_path}")
        except Exception as e:
            print(f"[agg] DOCX write skipped/failed: {e}")
    else:
        reason = 'skip flag' if skip_docx else f"row count {len(df)} > limit {max_docx_rows}"
        print(f"[agg] Skipping DOCX generation due to {reason}.")
    return paths


def compute_signature(secret: str, payload: bytes) -> str:
    return hmac.new(secret.encode('utf-8'), payload, hashlib.sha256).hexdigest()


def upload_results(upload_url: str, token: str, doc_id: str, request_id: str, paths: Dict[str,str], summary: Dict) -> None:
    if not upload_url:
        print('[agg] RESULT_UPLOAD_URL empty; skipping upload to app')
        return
    try:
        files = {}
        for key in ['csv','xlsx','docx']:
            p = paths.get(key)
            if p and os.path.exists(p):
                files[key] = (os.path.basename(p), open(p, 'rb'), 'application/octet-stream')
        data = { 'doc_id': doc_id or '', 'request_id': request_id or '', 'summary': json.dumps(summary) }
        headers = {}
        if token:
            headers['Authorization'] = f"Bearer {token}"
            headers['X-Extractor-Token'] = token
        resp = requests.post(upload_url, files=files, data=data, headers=headers, timeout=120)
        print('[agg] Upload results status:', resp.status_code)
    except Exception as e:
        print('[agg] Upload results error:', repr(e))


def main():
    chunks_dir = os.getenv('CHUNKS_DIR', 'chunks')
    original_filename = os.getenv('ORIGINAL_FILENAME', 'document.pdf')
    callback_url = os.getenv('CALLBACK_URL', '')
    callback_secret = os.getenv('CALLBACK_HMAC_SECRET', '')
    result_upload_url = os.getenv('RESULT_UPLOAD_URL', '')
    result_upload_token = os.getenv('RESULT_UPLOAD_TOKEN', '')
    doc_id = os.getenv('DOC_ID', '')
    request_id = os.getenv('REQUEST_ID', '').strip()
    callback_token = (os.getenv('CALLBACK_TOKEN', '').strip() or result_upload_token)
    page_start_env = os.getenv('PAGE_START', '').strip()
    page_end_env = os.getenv('PAGE_END', '').strip()
    pages_requested_env = os.getenv('PAGES_REQUESTED', '').strip()

    base = os.path.splitext(os.path.basename(original_filename))[0]
    pattern = os.path.join(chunks_dir, f"{base}-p*/*.csv")
    csv_files = sorted(glob.glob(pattern))
    if not csv_files:
        # also try flat files pattern
        csv_files = sorted(glob.glob(os.path.join(chunks_dir, f"**/{base}-p*.csv"), recursive=True))
    if not csv_files:
        raise SystemExit(f"No chunk CSV files found under {chunks_dir} for base {base}")

    print(f"[agg] Found {len(csv_files)} chunk CSV files to load")
    frames: List[pd.DataFrame] = []
    for idx, f in enumerate(csv_files, start=1):
        try:
            df = pd.read_csv(f)
            frames.append(df)
            if idx % 25 == 0 or idx == len(csv_files):
                print(f"[agg] Loaded {idx}/{len(csv_files)} files; last {os.path.basename(f)} -> {len(df)} rows")
        except Exception as e:
            print(f"[agg] Failed to read {f}: {e}")
    if not frames:
        raise SystemExit("No valid CSVs to aggregate")
    df_all = pd.concat(frames, ignore_index=True)
    # Normalize columns and deduplicate
    cols = ['surname','first_name','other_name','course_studied','faculty','grade','qualification_obtained','session']
    for c in cols:
        if c in df_all.columns:
            df_all[c] = df_all[c].astype(str).map(norm_space)
        else:
            df_all[c] = ''
    df_all = df_all[cols]
    before = len(df_all)
    df_all = df_all.drop_duplicates()
    after = len(df_all)
    print(f"[agg] Aggregated rows: {before} -> {after} (deduped)")

    out_dir = os.path.join('outputs')
    # DOCX controls via env
    skip_docx = (os.getenv('AGG_SKIP_DOCX', '').strip().lower() in ('1','true','yes'))
    try:
        max_docx_rows = int(os.getenv('AGG_MAX_DOCX_ROWS', '3000') or '3000')
    except Exception:
        max_docx_rows = 3000
    paths = save_outputs(df_all, base, out_dir, skip_docx=skip_docx, max_docx_rows=max_docx_rows)

    summary = {
        'status': 'success',
        'counts': {
            'rows': int(after),
            'pages_processed': int((int(page_end_env) - int(page_start_env) + 1) if page_start_env and page_end_env else (int(pages_requested_env) if pages_requested_env else 0)),
            'pages_with_results': int((int(page_end_env) - int(page_start_env) + 1) if page_start_env and page_end_env and after > 0 else ((int(pages_requested_env) if pages_requested_env else 0) if after > 0 else 0)),
        },
        'chunks': len(frames),
        'files': paths,
        'filename': original_filename,
        'doc_id': doc_id,
        'request_id': request_id,
    }

    # Upload to app
    if not result_upload_token:
        print('[agg] RESULT_UPLOAD_TOKEN is empty; upload may be rejected by app')
    upload_results(result_upload_url, result_upload_token, doc_id, request_id, paths, summary)

    # Callback to app
    if callback_url:
        payload = json.dumps(summary).encode('utf-8')
        headers = { 'Content-Type': 'application/json' }
        if callback_token:
            headers['Authorization'] = f"Bearer {callback_token}"
            headers['X-Extractor-Token'] = callback_token
        if callback_secret:
            headers['X-Extractor-Signature'] = compute_signature(callback_secret, payload)
        try:
            cr = requests.post(callback_url, data=payload, headers=headers, timeout=60)
            print('[agg] Callback status:', cr.status_code)
        except Exception as e:
            print('[agg] Callback error:', repr(e))
    else:
        print('[agg] CALLBACK_URL not set; skipping callback')

    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
